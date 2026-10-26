# -*- coding: utf-8 -*-
"""一页组会汇报 PPT（极简版）：纯白背景，简单列点工作内容。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

W, H = 13.333, 13.333 / 1.9          # 19:10 宽屏
prs = Presentation()
prs.slide_width = Inches(W); prs.slide_height = Inches(H)
slide = prs.slides.add_slide(prs.slide_layouts[6])

INK  = RGBColor(0x20, 0x20, 0x20)    # 标题
HEAD = RGBColor(0x17, 0x37, 0x5E)    # 小标题（深蓝，唯一一点色）
BODY = RGBColor(0x40, 0x40, 0x40)    # 正文
YAHEI = "微软雅黑"

def setfont(run, size, bold=False, color=INK):
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = color; run.font.name = YAHEI
    pr = run._r.get_or_add_rPr()
    ea = pr.find(qn('a:ea'))
    if ea is None:
        ea = etree.SubElement(pr, qn('a:ea'))
    ea.set('typeface', YAHEI)

def box(x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf

# ---- 标题 ----
tf = box(0.85, 0.55, 11.6, 0.62)
p = tf.paragraphs[0]
setfont(p.add_run(), 27, True, INK); p.runs[0].text = "近期工作汇报"

tf = box(0.85, 1.18, 11.6, 0.36)
p = tf.paragraphs[0]
setfont(p.add_run(), 13.5, False, RGBColor(0x80,0x80,0x80))
p.runs[0].text = "人机协作决策实验平台 · 部署 / 完善 / 测试"

# ---- 列点内容 ----
BLOCKS = [
    ("一、平台部署上线", [
        "阿里云 ECS 全链路部署：环境、数据库、前后端、Nginx、进程守护、定时备份",
        "公网可远程访问，数据库每日自动备份",
    ]),
    ("二、实验流程完善", [
        "问卷体系重构为数据库驱动：前置 20 题 + 后置 34 题，共 54 题、14 个构念",
        "补充反向计分题、情绪双极量表与分节指导语",
    ]),
    ("三、部署证明整理", [
        "6 大章节、11 张运行截图，覆盖 系统登录、后台核心功能、部署管理、行为日志",
    ]),
    ("四、全流程测试核查", [
        "17 个环节通过 16 项，覆盖 H / SOA / MOA 三组；定位问题 7 项",
    ]),
    ("下一步", [
        "优先修复邮件真实投递与单盲分组暴露两项高优先级问题，再进入正式被试招募",
    ]),
]

y = 1.60
for head, lines in BLOCKS:
    tf = box(0.85, y, 12.0, 0.34)
    p = tf.paragraphs[0]
    setfont(p.add_run(), 16, True, HEAD); p.runs[0].text = head
    y += 0.36
    tf = box(1.12, y, 11.7, 0.28 * len(lines) + 0.1)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2); p.line_spacing = 1.06
        setfont(p.add_run(), 13, False, BODY); p.runs[0].text = "•  " + ln
        y += 0.28
    y += 0.22

out = r"D:/27475/WorkBuddy-搭建网站/travel-experiment-platform/组会汇报_工作进展_一页.pptx"
prs.save(out)
print("SAVED:", out, " 比例 19:10")
