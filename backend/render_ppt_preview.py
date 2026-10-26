# -*- coding: utf-8 -*-
"""按 pptx 真实几何把幻灯片渲染成 PNG，用于排版核对（不依赖 Office）。"""
import os
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

PPTX = r"D:/27475/WorkBuddy-搭建网站/travel-experiment-platform/组会汇报_工作进展_一页.pptx"
OUT  = r"D:/27475/WorkBuddy-搭建网站/travel-experiment-platform/组会汇报_工作进展_一页_preview.png"
REG  = "C:/Windows/Fonts/msyh.ttc"
BLD  = "C:/Windows/Fonts/msyhbd.ttc"

prs = Presentation(PPTX)
sl = prs.slides[0]
SCALE = 96.0 / 914400.0          # EMU -> px (96 dpi)
Wpx = int(prs.slide_width * SCALE)
Hpx = int(prs.slide_height * SCALE)
img = Image.new("RGB", (Wpx, Hpx), "white")
d = ImageDraw.Draw(img)

def px(v): return int(v * SCALE)
def font(sz, bold=False):
    try:
        return ImageFont.truetype(BLD if bold else REG, int(sz * 96 / 72))
    except Exception:
        return ImageFont.load_default()

def rgb(c):
    try:    return "#%02X%02X%02X" % (c[0], c[1], c[2])
    except Exception: return None

def wrap(text, fnt, maxw):
    lines, cur = [], ""
    for ch in text:
        if d.textlength(cur + ch, font=fnt) <= maxw or not cur:
            cur += ch
        else:
            lines.append(cur); cur = ch
    if cur: lines.append(cur)
    return lines

for sh in sl.shapes:
    x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
    # 填充
    try:
        fill = sh.fill
        if fill.type is not None and fill.type == 1:  # solid
            col = rgb(fill.fore_color.rgb)
            if col:
                d.rectangle([x, y, x + w, y + h], fill=col)
    except Exception:
        pass
    # 描边
    try:
        if sh.line.fill.type == 1:
            lc = rgb(sh.line.color.rgb)
            lw = max(1, int((sh.line.width or 9525) * SCALE))
            if lc:
                d.rectangle([x, y, x + w, y + h], outline=lc, width=lw)
    except Exception:
        pass
    # 文本
    if sh.has_text_frame and sh.text_frame.text.strip():
        tf = sh.text_frame
        # 垂直居中近似
        total = 0
        blocks = []
        for para in tf.paragraphs:
            txt = "".join(r.text for r in para.runs)
            if not txt:
                blocks.append((0, None, txt, [])); continue
            sz = para.runs[0].font.size.pt if para.runs[0].font.size else 18
            bold = bool(para.runs[0].font.bold)
            fnt = font(sz, bold)
            maxw = w - 6
            ls = wrap(txt, fnt, maxw)
            lh = int(sz * 96 / 72 * 1.22)
            blocks.append((lh, (fnt, rgb(para.runs[0].font.color.rgb) or "#000"),
                           txt, ls))
            total += lh * len(ls) + 4
        cy = y + max(0, (h - total) // 2) if (h and total and total < h) else y
        for lh, meta, txt, ls in blocks:
            if not meta:
                cy += lh or 6; continue
            fnt, col = meta
            for line in ls:
                d.text((x + 3, cy), line, font=fnt, fill=col)
                cy += lh

img.save(OUT)
print("SAVED:", OUT, img.size)
